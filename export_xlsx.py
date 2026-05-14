from pathlib import Path
from typing import List, Dict, Any
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


HEADER_FONT = Font(name='Arial', bold=True, size=11, color='FFFFFF')
HEADER_FILL = PatternFill('solid', fgColor='2F5496')
HEADER_ALIGN = Alignment(horizontal='center', vertical='center', wrap_text=True)

DATA_FONT = Font(name='Arial', size=10)
LEVEL_FILL = PatternFill('solid', fgColor='D6E4F0')
THIN_BORDER = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin'),
)


def _apply_header(ws, row, cols):
    for col_idx, title in enumerate(cols, 1):
        cell = ws.cell(row=row, column=col_idx, value=title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
        cell.border = THIN_BORDER


def _apply_data(ws, row, values):
    for col_idx, val in enumerate(values, 1):
        cell = ws.cell(row=row, column=col_idx, value=val)
        cell.font = DATA_FONT
        cell.border = THIN_BORDER
        cell.alignment = Alignment(wrap_text=True)


def _auto_width(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 4, 60)


def export_all(micro_results: List[Dict],
               meso_results: List[Dict],
               macro_result: Dict,
               out_path: Path):
    wb = Workbook()

    _write_micro(wb, micro_results)
    _write_meso(wb, meso_results)
    _write_macro(wb, macro_result)

    if 'Sheet' in wb.sheetnames:
        del wb['Sheet']

    wb.save(out_path)


def _write_micro(wb: Workbook, results: List[Dict]):
    ws = wb.create_sheet('MICRO')
    headers = ['Robot ID', 'Status', 'Level', 'Solution',
               'Variable', 'Equation', 'LaTeX']
    _apply_header(ws, 1, headers)

    row = 2
    for res in results:
        rid = res.get('robot_id', '?')
        status = res.get('status', '?')

        if not res.get('equations'):
            _apply_data(ws, row, [rid, status, '', '', '', '', ''])
            row += 1
            continue

        for eq in res['equations']:
            _apply_data(ws, row, [
                rid,
                status,
                eq.get('level', ''),
                eq.get('solution', ''),
                eq.get('variable', ''),
                eq.get('text_form', ''),
                eq.get('latex_form', ''),
            ])
            row += 1

    _auto_width(ws)


def _write_meso(wb: Workbook, results: List[Dict]):
    ws = wb.create_sheet('MESO')
    headers = ['Robot Group', 'Status', 'Level', 'Solution',
               'Variable', 'Equation', 'LaTeX']
    _apply_header(ws, 1, headers)

    row = 2
    for res in results:
        group_str = ', '.join(str(r) for r in res.get('robot_ids', []))
        status = res.get('status', '?')

        if not res.get('equations'):
            _apply_data(ws, row, [group_str, status, '', '', '', '', ''])
            row += 1
            continue

        for eq in res['equations']:
            _apply_data(ws, row, [
                group_str,
                status,
                eq.get('level', ''),
                eq.get('solution', ''),
                eq.get('variable', ''),
                eq.get('text_form', ''),
                eq.get('latex_form', ''),
            ])
            row += 1

    _auto_width(ws)

    if results and results[0].get('coeff_table') is not None:
        df = results[0]['coeff_table']
        if not df.empty:
            ws_coeff = wb.create_sheet('MESO_coefficients')
            _apply_header(ws_coeff, 1, ['index'] + list(df.columns))
            for i, (idx, row_data) in enumerate(df.iterrows()):
                values = [idx] + [round(v, 8) if isinstance(v, float) else v
                                  for v in row_data.values]
                _apply_data(ws_coeff, i + 2, values)
            _auto_width(ws_coeff)


def _write_macro(wb: Workbook, result: Dict):
    ws = wb.create_sheet('MACRO')
    headers = ['Robot Idx', 'Status', 'Level', 'Solution',
               'Variable', 'Equation', 'LaTeX']
    _apply_header(ws, 1, headers)

    row = 2
    ridx = result.get('robot_idx', '?')
    status = result.get('status', '?')

    if not result.get('equations'):
        _apply_data(ws, row, [ridx, status, '', '', '', '', ''])
        return

    for eq in result['equations']:
        _apply_data(ws, row, [
            ridx,
            status,
            eq.get('level', ''),
            eq.get('solution', ''),
            eq.get('variable', ''),
            eq.get('text_form', ''),
            eq.get('latex_form', ''),
        ])
        row += 1

    _auto_width(ws)


def export_micro_only(micro_results: List[Dict], out_path: Path):
    wb = Workbook()
    _write_micro(wb, micro_results)
    if 'Sheet' in wb.sheetnames:
        del wb['Sheet']
    wb.save(out_path)


def export_meso_only(meso_results: List[Dict], out_path: Path):
    wb = Workbook()
    _write_meso(wb, meso_results)
    if 'Sheet' in wb.sheetnames:
        del wb['Sheet']
    wb.save(out_path)


def export_macro_only(macro_result: Dict, out_path: Path):
    wb = Workbook()
    _write_macro(wb, macro_result)
    if 'Sheet' in wb.sheetnames:
        del wb['Sheet']
    wb.save(out_path)
